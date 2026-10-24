import json
import os
from pathlib import Path
import tempfile
import unittest

import openpyxl

from build_extended_newborn import count, extract


class ExtendedNewbornTests(unittest.TestCase):
    def test_missing_counts_are_not_zero(self):
        for value in (None, '', '#DIV/0!', -1, 1.5, True):
            with self.assertRaises(ValueError):
                count(value, 'test cell')
        self.assertEqual(count(0, 'test cell'), 0)

    def fixture(self, directory):
        path = Path(directory) / 'fixture.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Sheet1'
        headers = {'F8': '8th of Feb -31st of Augast 2026', 'F9': 'Hospital/ PC Name',
                   'G9': 'Deliveries in the institution', 'J9': 'Screened',
                   'M9': 'Recived by the labs', 'N9': 'Rejected by the labs', 'O9': 'Coverage % of Omani'}
        for cell, value in headers.items():
            sheet[cell] = value
        for column, value in enumerate(['Test Hospital', 10, 2, 12, 11, 2, 13, 15, 1, 110], 6):
            sheet.cell(11, column, value)
        book.save(path)
        return path

    def test_over_100_ratio_is_retained(self):
        with tempfile.TemporaryDirectory() as directory:
            record = extract(self.fixture(directory))['institutions'][0]
            self.assertEqual(record['screened'], 13)
            self.assertAlmostEqual(record['omaniScreeningRatio'], 110)

    def test_bad_totals_and_missing_counts_fail(self):
        for cell, value in [('I11', 99), ('G11', None), ('N11', 20), ('O11', 90)]:
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as directory:
                path = self.fixture(directory)
                book = openpyxl.load_workbook(path)
                book.active[cell] = value
                book.save(path)
                with self.assertRaises(ValueError):
                    extract(path)

    def test_duplicate_institution_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.fixture(directory)
            book = openpyxl.load_workbook(path)
            for column in range(6, 16):
                book.active.cell(12, column, book.active.cell(11, column).value)
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'duplicate institution'):
                extract(path)

    def test_changed_period_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = self.fixture(directory)
            book = openpyxl.load_workbook(path)
            book.active['F8'] = 'September 2026'
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'Reporting period changed'):
                extract(path)

    @unittest.skipUnless(os.environ.get('ENS_SOURCE'), 'Set ENS_SOURCE for source verification')
    def test_original_workbook_matches_import(self):
        saved = json.loads((Path(__file__).resolve().parents[1] / 'src/data/extendedNewborn2026.json').read_text())
        self.assertEqual(extract(Path(os.environ['ENS_SOURCE'])), saved)
        self.assertEqual(len(saved['institutions']), 5)
        self.assertEqual(saved['period'], {'start': '2026-02-08', 'end': '2026-08-31'})
        for key, expected in {'deliveries': 4748, 'screened': 4092, 'received': 4785, 'rejected': 132,
                              'omaniDeliveries': 4711, 'nonOmaniDeliveries': 37,
                              'omaniScreened': 4087, 'nonOmaniScreened': 5}.items():
            self.assertEqual(sum(row[key] for row in saved['institutions']), expected)


if __name__ == '__main__':
    unittest.main()
