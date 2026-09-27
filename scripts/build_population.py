#!/usr/bin/env python3
"""
Build src/data/population.json from the NBG 2025 population-estimates workbook
(التقديرات السكانية 2025 — one sheet per wilayah, one block per health institution).

Usage:
    pip install openpyxl
    python3 scripts/build_population.py path/to/population_2025.xlsx

Workbook layout (as received, July 2026):
  • inst   — summary: Omani / expatriate / total per institution (29 rows)
  • فئات   — department target-group ratios (flat governorate %, not used here)
  • صحار, صحم, الخابوره, السويق, شناص, لوى — per-institution blocks:
      title row (institution name) → 5 header rows → 17 five-year age bands
      (0-4 … 80+) → الجملة (total) row.  Columns: B/C = Omani M/F,
      G/H = non-Omani M/F.

The script fails loudly (non-zero exit) if any block's age bands do not sum to
its own total row, so a malformed sheet can never silently reach the app.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "src" / "data" / "population.json"

AGE_BANDS = [
    "0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39",
    "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79", "80+",
]

# Sheet name → (English wilayah name used across the app, Arabic name)
WILAYAT = [
    ("صحار", "Sohar", "صحار"),
    ("صحم", "Saham", "صحم"),
    ("السويق", "As Suwayq", "السويق"),
    ("الخابوره", "Al Khabourah", "الخابورة"),
    ("شناص", "Shinas", "شناص"),
    ("لوى", "Liwa", "لوى"),
]

# Arabic institution name (as written in the wilayah sheets) → clean English name.
# English names follow the workbook's own `inst` sheet, normalised for display.
INSTITUTION_EN = {
    "مجمع صحار الصحي": "Sohar EHC",
    "مركز الملتقى الصحي": "Al Multaqa HC",
    "مستشفى وادي حيبي": "Wadi Hibi Hospital",
    "مركز وادي عاهن الصحي": "Wadi Ahin HC",
    "مركز العوينات الصحي": "Al Uwainat HC",
    "مركز فلج القبائل الصحي": "Falaj Al Qabail HC",
    "مركز الطريف الصحي": "Al Traif HC",
    "مجمع صحم الصحي": "Saham EHC",
    "مركز وادي بني عمر الصحي": "Wadi Bani Umar HC",
    "مركز حفيت الصحي": "Hafeet HC",
    "مركز الغويصة الصحي": "Al Ghwaisah HC",
    "مجمع الخابورة الصحي": "Al Khabourah EHC",
    "مركز وادي شافان الصحي": "Wadi Shafan HC",
    "مستشفى وادي الصرمي": "Wadi As Sarami Hospital",
    "مستشفى وادي الحواسنة": "Wadi Al Hawasinah Hospital",
    "مركز القصبية الصحي": "Al Qasabiyah HC",
    "مجمع السويق الصحي": "As Suwayq EHC",
    "مركز البداية الصحي": "Al Bidayah HC",
    "مركز الخضراء الصحي": "Al Khadra HC",
    "مركز وادي الجهاور الصحي": "Wadi Al Jahawir HC",
    "مركز المشايق الصحي": "Al Mashaiq HC",
    "مركز الشريسة الصحي": "Al Shreesah HC",
    "مركز الثرمد الصحي": "Al Tharmed HC",
    "مجمع شناص الصحي": "Shinas EHC",
    "مركز ابو بقرة الصحي": "Abu Baqra HC",
    "مركز سور البلوش الصحي (جديد)": "Sur Al Balush HC (new)",
    "مركز لوى الصحي": "Liwa HC",
    "مركز نبر الصحي": "Nabur HC",
    "مركز رحب الصحي": "Rahab HC",
}


def norm_ar(name: str) -> str:
    """Normalise Arabic spelling differences between the summary and detail sheets."""
    return (
        name.strip()
        .replace("أ", "ا")
        .replace("مجمع لوى", "مركز لوى")
        .replace("مركز مشايق", "مركز المشايق")
    )


def r2(v: float) -> float:
    """Two decimals is the least precision at which every total still reconciles."""
    return round(float(v), 2)


def fail(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def read_blocks(ws) -> list[dict]:
    blocks = []
    r = 1
    while r <= ws.max_row:
        title = ws.cell(r, 1).value
        marker = ws.cell(r + 1, 2).value
        if isinstance(title, str) and marker and "عماني" in str(marker):
            first = r + 6
            labels = [str(ws.cell(first + i, 1).value).replace(" ", "") for i in range(17)]
            # Source writes bands right-to-left ("4 - 0"); compare as sets of bounds.
            for got, want in zip(labels, AGE_BANDS):
                if sorted(got.split("-")) != sorted(want.split("-")):
                    fail(f"{ws.title} / {title}: unexpected age band '{got}' (wanted '{want}')")
            cols = {"om_m": 2, "om_f": 3, "ex_m": 7, "ex_f": 8}
            data = {k: [ws.cell(first + i, c).value or 0 for i in range(17)] for k, c in cols.items()}
            total_row = first + 17
            if str(ws.cell(total_row, 1).value).strip() != "الجملة":
                fail(f"{ws.title} / {title}: total row not found at row {total_row}")
            for k, c in cols.items():
                src_total = ws.cell(total_row, c).value or 0
                if abs(sum(data[k]) - src_total) > 0.5:
                    fail(f"{ws.title} / {title}: {k} bands sum {sum(data[k]):.1f} ≠ total {src_total:.1f}")
            blocks.append({"ar": title.strip(), **data})
            r = total_row + 1
        else:
            r += 1
    return blocks


def main() -> None:
    if len(sys.argv) != 2:
        fail("usage: build_population.py <workbook.xlsx>")
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

    # ---- summary sheet (for reconciliation only) ----------------------------
    inst = wb["inst"]
    summary = {}
    for row in inst.iter_rows(min_row=3, values_only=True):
        name, om, ex = row[0], row[1], row[2]
        if isinstance(name, str) and isinstance(om, (int, float)):
            summary[norm_ar(name)] = (om, ex)
    summary_totals = {
        "institutions": len(summary),
        "omani": sum(v[0] for v in summary.values()),
        "expatriate": sum(v[1] for v in summary.values()),
    }

    # ---- detail sheets -------------------------------------------------------
    institutions = []
    for sheet, wil_en, _wil_ar in WILAYAT:
        for b in read_blocks(wb[sheet]):
            en = INSTITUTION_EN.get(b["ar"])
            if not en:
                fail(f"no English name for institution '{b['ar']}' in sheet {sheet}")
            om, ex = sum(b["om_m"]) + sum(b["om_f"]), sum(b["ex_m"]) + sum(b["ex_f"])
            key = norm_ar(b["ar"])
            in_summary = key in summary
            if in_summary:
                s_om, s_ex = summary[key]
                if abs(om - s_om) > 1 or abs(ex - s_ex) > 1:
                    fail(f"{b['ar']}: detail {om:.0f}/{ex:.0f} ≠ summary {s_om}/{s_ex}")
            institutions.append(
                {
                    "en": en,
                    "ar": b["ar"].replace(" (جديد)", ""),
                    "wilayat": wil_en,
                    "inSummary": in_summary,
                    "omM": [r2(v) for v in b["om_m"]],
                    "omF": [r2(v) for v in b["om_f"]],
                    "exM": [r2(v) for v in b["ex_m"]],
                    "exF": [r2(v) for v in b["ex_f"]],
                }
            )

    detail_keys = {norm_ar(i["ar"]) for i in institutions}
    missing = sorted(set(summary) - detail_keys)
    if missing:
        fail(f"summary institutions with no detail block: {missing}")

    out = {
        "meta": {
            "title": "Population Estimates 2025",
            "titleAr": "التقديرات السكانية لعام 2025",
            "source": "NBG population estimates 2025 workbook (MOH, updated July 2026)",
            "reference": "End-year 2025 estimates by health-institution catchment",
            "ageBands": AGE_BANDS,
            "summarySheet": summary_totals,
        },
        "wilayats": [{"en": en, "ar": ar} for _s, en, ar in WILAYAT],
        "institutions": institutions,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

    om = sum(sum(i["omM"]) + sum(i["omF"]) for i in institutions)
    ex = sum(sum(i["exM"]) + sum(i["exF"]) for i in institutions)
    print(f"Wrote {OUT.relative_to(OUT.parents[2])}: {len(institutions)} institutions, "
          f"Omani {om:,.0f}, expatriate {ex:,.0f}, total {om + ex:,.0f}")
    extra = [i["en"] for i in institutions if not i["inSummary"]]
    if extra:
        print(f"NOTE: in detail sheets but not in the summary sheet: {extra}")


if __name__ == "__main__":
    main()
