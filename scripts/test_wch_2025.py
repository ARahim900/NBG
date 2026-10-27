import json
import os
from pathlib import Path
import tempfile
import unittest

import openpyxl

from build_wch_2025 import (DS_CENTRES, VISIT_CENTRES, count, extract_down_syndrome, extract_maltreatment,
                            extract_visits)

DATA = Path(__file__).resolve().parents[1] / 'src' / 'data'


def visit_sheet(book, title, header_row=7):
    sheet = book.create_sheet(title)
    sheet.cell(header_row, 2, 'Institution Name / Months')
    for m in range(12):
        for i, label in enumerate(('Target ', 'Clerked ', 'Refarred ')):
            sheet.cell(header_row + 1, 3 + m * 3 + i, label)
    sheet.cell(header_row, 40, 'Total ')
    row = header_row + 2
    for name in VISIT_CENTRES:
        sheet.cell(row, 2, name.title())
        for m in range(12):
            sheet.cell(row, 3 + m * 3, 10)
            sheet.cell(row, 4 + m * 3, 2)
            sheet.cell(row, 5 + m * 3, 0)
        sheet.cell(row, 40, 30)
        sheet.cell(row, 41, 24)
        sheet.cell(row, 42, 0)
        row += 1
    sheet.cell(row, 2, 'Total per month ')
    sheet.cell(row, 40, 30 * len(VISIT_CENTRES))
    sheet.cell(row, 41, 24 * len(VISIT_CENTRES))
    sheet.cell(row, 42, 0)
    return sheet


def visits_fixture(directory):
    book = openpyxl.Workbook()
    book.remove(book.active)
    visit_sheet(book, '3 years Visits 2025')
    visit_sheet(book, '4 years Visits 2025', header_row=8)
    path = Path(directory) / 'visits.xlsx'
    book.save(path)
    return path


def ds_fixture(directory):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A2'] = 'Governorate: North Al Batina Year: January–December 2025'
    for cell, value in {'A4': 'Institution Name', 'G4': '***Number of newly diagnosed cases', 'L6': 'Male',
                        'M6': 'Female', 'A18': 'Total number of cases with comorbidities ', 'B18': 'Cardiac'}.items():
        sheet[cell] = value
    for i, (reg, morb, _, _) in enumerate(DS_CENTRES):
        sheet.cell(7 + i, 1, reg.title())
        for col, value in zip(range(7, 14), (1, 10, 5, 4, 1, 6, 4)):
            sheet.cell(7 + i, col, value)
        sheet.cell(23 + i, 1, morb)
        for col in range(2, 15):
            sheet.cell(23 + i, col, 1)
        sheet.cell(23 + i, 10, 8)  # normal weight → 1 + 8 + 1 = 10 children
    sheet['A15'] = 'Total'
    sheet['B15'], sheet['C15'] = 6, 6
    for col, value in zip(range(7, 14), (1, 10, 5, 4, 1, 6, 4)):
        sheet.cell(15, col, value * len(DS_CENTRES))
    path = Path(directory) / 'ds.xlsx'
    book.save(path)
    return path


def mt_fixture(directory, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(['Reporting Institute', 'Reporting Region', 'Reporting Wilayat', 'Note Date', 'Age', 'Sex',
                  'Nationality', 'Maltreatment Type'])
    for row in rows:
        sheet.append(row)
    path = Path(directory) / 'mt.xlsx'
    book.save(path)
    return path


GOOD_NOTICE = ['Sohar Hospital', 'NORTH BATINAH', 'Sohar', '08-01-2025', 4, 'F', 'SULTANATE OF OMAN', 'Neglect']


class CountTests(unittest.TestCase):
    def test_blank_is_null_only_when_allowed(self):
        self.assertIsNone(count(None, 'x', allow_blank=True))
        self.assertEqual(count(0, 'x'), 0)
        for value in (None, '', -1, 1.5, True, '#DIV/0!'):
            with self.assertRaises(ValueError):
                count(value, 'x')


class VisitTests(unittest.TestCase):
    def test_reads_every_centre(self):
        with tempfile.TemporaryDirectory() as d:
            data = extract_visits(visits_fixture(d))
            self.assertEqual(len(data['three']['centres']), len(VISIT_CENTRES))
            self.assertEqual(data['four']['statedTotals']['clerked'], 24 * len(VISIT_CENTRES))
            self.assertEqual(data['three']['flags'], [])

    def test_blank_month_stays_null_and_is_flagged(self):
        with tempfile.TemporaryDirectory() as d:
            path = visits_fixture(d)
            book = openpyxl.load_workbook(path)
            for col in (3, 4, 5):
                book['3 years Visits 2025'].cell(9, col).value = None
            book.save(path)
            data = extract_visits(path)
            self.assertIsNone(data['three']['centres'][0]['monthly'][0])
            self.assertTrue(any('no figures for Jan' in f for f in data['three']['flags']))

    def test_total_that_disagrees_with_months_is_kept_and_flagged(self):
        with tempfile.TemporaryDirectory() as d:
            path = visits_fixture(d)
            book = openpyxl.load_workbook(path)
            book['3 years Visits 2025'].cell(9, 5).value = 30  # 30 referrals from 2 visits in Jan
            book.save(path)
            data = extract_visits(path)
            self.assertEqual(data['three']['centres'][0]['referred'], 0)
            flags = ' '.join(data['three']['flags'])
            self.assertIn('30 referrals from 2 visits', flags)
            self.assertIn('referred months add to 30', flags)

    def test_unknown_institution_fails(self):
        with tempfile.TemporaryDirectory() as d:
            path = visits_fixture(d)
            book = openpyxl.load_workbook(path)
            book['3 years Visits 2025'].cell(9, 2).value = 'Somewhere Else'
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'unknown institution'):
                extract_visits(path)


class DownSyndromeTests(unittest.TestCase):
    def test_reads_clinics(self):
        with tempfile.TemporaryDirectory() as d:
            data = extract_down_syndrome(ds_fixture(d))
            self.assertEqual(len(data['centres']), 6)
            self.assertEqual(data['centres'][0]['comorbidities']['cardiac'], 1)
            self.assertEqual(data['flags'], [])

    def test_gender_must_match_registered(self):
        with tempfile.TemporaryDirectory() as d:
            path = ds_fixture(d)
            book = openpyxl.load_workbook(path)
            book.active['L7'] = 99
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'males \\+ females'):
                extract_down_syndrome(path)

    def test_age_mismatch_is_flagged_not_fixed(self):
        with tempfile.TemporaryDirectory() as d:
            path = ds_fixture(d)
            book = openpyxl.load_workbook(path)
            book.active['I7'] = 6
            book.active['I15'] = 31
            book.save(path)
            data = extract_down_syndrome(path)
            self.assertEqual(data['centres'][0]['age0to5'], 6)
            self.assertIn('age groups add to 11', data['flags'][0])

    def test_other_year_fails(self):
        with tempfile.TemporaryDirectory() as d:
            path = ds_fixture(d)
            book = openpyxl.load_workbook(path)
            book.active['A2'] = 'Year: January–December 2026'
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'Reporting period'):
                extract_down_syndrome(path)


class MaltreatmentTests(unittest.TestCase):
    def test_aggregates_only(self):
        with tempfile.TemporaryDirectory() as d:
            data = extract_maltreatment(mt_fixture(d, [GOOD_NOTICE, GOOD_NOTICE]))
            self.assertEqual(data['total'], 2)
            self.assertEqual(data['byMonth'][0], 2)
            self.assertEqual(next(w for w in data['byWilayat'] if w['wilayat'] == 'Sohar')['Neglect'], 2)
            self.assertNotIn('records', data)
            self.assertNotIn('08-01-2025', json.dumps(data))

    def test_out_of_year_or_unknown_values_fail(self):
        for index, value in [(3, '08-01-2024'), (2, 'Muscat'), (7, 'Bullying'), (4, 19)]:
            row = list(GOOD_NOTICE)
            row[index] = value
            with self.subTest(value=value), tempfile.TemporaryDirectory() as d:
                with self.assertRaises(ValueError):
                    extract_maltreatment(mt_fixture(d, [row]))


@unittest.skipUnless(os.environ.get('WCH_SOURCE_DIR'), 'Set WCH_SOURCE_DIR to verify against the source workbooks')
class SourceTests(unittest.TestCase):
    def test_saved_json_matches_workbooks(self):
        src = Path(os.environ['WCH_SOURCE_DIR'])
        self.assertEqual(extract_visits(src / '3 & 4 years Visits 2025.xlsx'),
                         json.loads((DATA / 'childVisits2025.json').read_text()))
        self.assertEqual(extract_down_syndrome(src / 'Suhar 2025 Down syndrome .xlsx'),
                         json.loads((DATA / 'downSyndrome2025.json').read_text()))
        self.assertEqual(extract_maltreatment(src / 'Maltreatment Notifications 2025.xlsx'),
                         json.loads((DATA / 'maltreatment2025.json').read_text()))

    def test_headline_figures(self):
        visits = json.loads((DATA / 'childVisits2025.json').read_text())
        self.assertEqual(visits['three']['statedTotals'], {'target': 14591, 'clerked': 5341, 'referred': 817})
        self.assertEqual(visits['four']['statedTotals'], {'target': 14666, 'clerked': 5244, 'referred': 211})
        ds = json.loads((DATA / 'downSyndrome2025.json').read_text())
        self.assertEqual(sum(c['registered'] for c in ds['centres']), 295)
        self.assertEqual(sum(c['newCases'] for c in ds['centres']), 27)
        self.assertEqual(sum(c['age6to12'] for c in ds['centres']), 116)
        mt = json.loads((DATA / 'maltreatment2025.json').read_text())
        self.assertEqual(mt['total'], 204)


if __name__ == '__main__':
    unittest.main()
