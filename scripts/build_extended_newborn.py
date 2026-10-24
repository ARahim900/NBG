"""Read the supplied February–August ENS summary without changing the workbook."""
import argparse
import hashlib
import json
from pathlib import Path

import openpyxl


def count(value, location):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0 or value != int(value):
        raise ValueError(f'Invalid or missing count at {location}: {value!r}')
    return int(value)


def extract(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    if workbook.sheetnames != ['Sheet1']:
        raise ValueError('Unexpected workbook sheets; review the source layout before importing.')
    sheet = workbook['Sheet1']
    if sheet['F8'].value != '8th of Feb -31st of Augast 2026':
        raise ValueError('Reporting period changed; review the source period before importing.')
    expected_headers = {'F9': 'Hospital/ PC Name', 'G9': 'Deliveries in the institution',
                        'J9': 'Screened', 'M9': 'Recived by the labs',
                        'N9': 'Rejected by the labs', 'O9': 'Coverage % of Omani'}
    for cell, expected in expected_headers.items():
        if str(sheet[cell].value).strip() != expected:
            raise ValueError(f'Unexpected heading at {cell}')
    institutions = []
    seen = set()
    for row in range(11, sheet.max_row + 1):
        if all(sheet.cell(row, col).value is None for col in range(6, 16)):
            continue
        name = str(sheet.cell(row, 6).value or '').strip()
        if not name or name.casefold() in seen:
            raise ValueError(f'Missing or duplicate institution at F{row}')
        seen.add(name.casefold())
        record = {'sourceRow': row, 'sourceName': name,
                  'name': name.replace('hospital', 'Hospital').replace('Polycinic', 'Polyclinic')}
        for col, key in [(7, 'omaniDeliveries'), (8, 'nonOmaniDeliveries'), (9, 'deliveries'),
                         (10, 'omaniScreened'), (11, 'nonOmaniScreened'), (12, 'screened'),
                         (13, 'received'), (14, 'rejected')]:
            record[key] = count(sheet.cell(row, col).value, f'Sheet1 row {row}, column {col}')
        for total, parts in [('deliveries', ('omaniDeliveries', 'nonOmaniDeliveries')),
                             ('screened', ('omaniScreened', 'nonOmaniScreened'))]:
            if record[total] != sum(record[key] for key in parts):
                raise ValueError(f'{total} total does not reconcile at row {row}')
        if record['rejected'] > record['received']:
            raise ValueError(f'Rejected exceeds received at row {row}')
        coverage = sheet.cell(row, 15).value
        ratio = record['omaniScreened'] / record['omaniDeliveries'] * 100 if record['omaniDeliveries'] else None
        if ratio is not None and (not isinstance(coverage, (int, float)) or abs(coverage - ratio) > 0.000001):
            raise ValueError(f'Omani percentage does not reconcile at O{row}')
        record['omaniScreeningRatio'] = ratio
        institutions.append(record)
    if not institutions:
        raise ValueError('No institution data found')
    return {'source': {'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                       'sheet': 'Sheet1', 'range': f'F9:O{sheet.max_row}', 'periodLabel': sheet['F8'].value},
            'period': {'start': '2026-02-08', 'end': '2026-08-31'}, 'institutions': institutions}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    args.output.write_text(json.dumps(extract(args.source), indent=2) + '\n')
